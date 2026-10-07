"""core.py - everything except the UI.

Pipeline per (surah, language):
  pair files -> Whisper word timestamps (recitation + interpretation)
  -> verse starts: difflib against the reference docx (first pass), Gemini checks/corrects (second pass)
  -> measured verse spans (snapped to silence) -> <tag>.verses.json -> combined mp3
"""
from __future__ import annotations

import base64
import difflib
import hashlib
import io
import json
import re
import time
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np

VERSE_COUNTS = [7, 286, 200, 176, 120, 165, 206, 75, 129, 109, 123, 111, 43, 52, 99, 128, 111, 110, 98, 135, 112,
                78, 118, 64, 77, 227, 93, 88, 69, 60, 34, 30, 73, 54, 45, 83, 182, 88, 75, 85, 54, 53, 89, 59, 37,
                35, 38, 29, 18, 45, 60, 49, 62, 55, 78, 96, 29, 22, 24, 13, 14, 11, 11, 18, 12, 12, 30, 52, 52, 44,
                28, 28, 20, 56, 40, 31, 50, 40, 46, 42, 29, 19, 36, 25, 22, 17, 19, 26, 30, 20, 15, 21, 11, 8, 8,
                19, 5, 8, 8, 11, 11, 8, 3, 9, 5, 4, 7, 3, 6, 3, 5, 4, 5, 6]
SURAH_NAMES = ["الفاتحة", "البقرة", "آل عمران", "النساء", "المائدة", "الأنعام", "الأعراف", "الأنفال", "التوبة", "يونس", "هود", "يوسف", "الرعد", "إبراهيم", "الحجر", "النحل", "الإسراء", "الكهف", "مريم", "طه", "الأنبياء", "الحج", "المؤمنون", "النور", "الفرقان", "الشعراء", "النمل", "القصص", "العنكبوت", "الروم", "لقمان", "السجدة", "الأحزاب", "سبأ", "فاطر", "يس", "الصافات", "ص", "الزمر", "غافر", "فصلت", "الشورى", "الزخرف", "الدخان", "الجاثية", "الأحقاف", "محمد", "الفتح", "الحجرات", "ق", "الذاريات", "الطور", "النجم", "القمر", "الرحمن", "الواقعة", "الحديد", "المجادلة", "الحشر", "الممتحنة", "الصف", "الجمعة", "المنافقون", "التغابن", "الطلاق", "التحريم", "الملك", "القلم", "الحاقة", "المعارج", "نوح", "الجن", "المزمل", "المدثر", "القيامة", "الإنسان", "المرسلات", "النبأ", "النازعات", "عبس", "التكوير", "الانفطار", "المطففين", "الانشقاق", "البروج", "الطارق", "الأعلى", "الغاشية", "الفجر", "البلد", "الشمس", "الليل", "الضحى", "الشرح", "التين", "العلق", "القدر", "البينة", "الزلزلة", "العاديات", "القارعة", "التكاثر", "العصر", "الهمزة", "الفيل", "قريش", "الماعون", "الكوثر", "الكافرون", "النصر", "المسد", "الإخلاص", "الفلق", "الناس"]
AUDIO_EXT = {".mp3", ".m4a", ".wav", ".flac", ".ogg", ".opus", ".aac", ".webm", ".mp4", ".wma"}
REF_EXT = {".docx", ".txt"}
RATE = 24000
OUT_RATE = 24000

LANG_CODES = {"spanish": "es", "french": "fr", "persian": "fa", "farsi": "fa", "english": "en", "urdu": "ur",
              "turkish": "tr", "indonesian": "id", "german": "de", "russian": "ru", "portuguese": "pt",
              "brazilian": "pt", "bengali": "bn", "hindi": "hi", "malay": "ms", "chinese": "zh",
              "italian": "it", "dutch": "nl", "swahili": "sw", "hausa": "ha", "pashto": "ps", "kurdish": "ku"}


def lang_code(lang: str) -> str:
    """Language folder name -> Whisper code. Anything unknown (Egyptian, Saudi, Moroccan...) is Arabic."""
    return LANG_CODES.get(lang.split()[0].lower(), "ar") if lang.strip() else "ar"


@dataclass
class Settings:
    taf_dir: Path | None = None
    rec_dir: Path | None = None
    ref_dir: Path | None = None
    final_dir: Path = Path("final")
    api_key: str = ""
    model: str = "gemini-3.1-pro-preview"
    whisper_model: str = "large-v3"
    surah_from: int = 1
    surah_to: int = 114
    languages: tuple = ()
    gap_after_rec_ms: int = 400
    gap_after_taf_ms: int = 900
    bitrate: str = "64k"
    window: int = 20              # verses per Gemini request
    disagree_words: int = 4       # Gemini vs the document-based guess: more than this -> flag the verse
    min_score: float = 0.40       # head/tail text match below this -> flag the verse
    overwrite: bool = False
    snap_ms: int = 400            # how far (ms) from Whisper's word time a cut may move to reach a real pause
    timeout_s: int = 150          # give up on one Gemini request after this long
    usage: dict = field(default_factory=lambda: {"calls": 0, "in": 0, "out": 0, "audio_s": 0.0})

    @property
    def work_dir(self) -> Path:
        return Path(self.final_dir) / "_work"


@dataclass
class Job:
    number: int
    lang: str
    name: str
    taf: Path
    rec: Path
    ref: Path | None = None
    units: list | None = None      # [(first verse, last verse)]: verses explained together are ONE unit

    @property
    def tag(self):
        return f"{self.number:03d}_{self.lang}"


@dataclass
class Word:
    s: int
    e: int
    t: str


# ------------------------------------------------------------------------------ pairing ------- #
TAF_RE = re.compile(r"^\s*([^_]+?)_(\d{1,3})\s*([^_]*)")      # Egyptian_100 سورة العاديات_Tafseer_MM
REC_RE = re.compile(r"^\s*(\d{1,3})(?!\d)")                    # 100 - الشيخ ... ｜ سورة العاديات ｜ ...


def _files(folder, exts):
    return sorted(p for p in Path(folder).rglob("*")
                  if p.is_file() and p.suffix.lower() in exts and not p.name.startswith(("~$", "._")))


def pair(taf_dir, rec_dir, ref_dir=None):
    """-> (jobs, problems). Pairing key = surah number (recitation) and (language, number) (interpretation)."""
    problems, taf, rec, ref = [], {}, {}, {}
    for p in _files(taf_dir, AUDIO_EXT):
        m = TAF_RE.match(p.stem)
        if not m or not 1 <= int(m[2]) <= 114:
            problems.append(f"interpretation file not recognised: {p.name}")
        else:
            taf.setdefault((m[1].strip(), int(m[2])), (p, m[3].strip()))
    for p in _files(rec_dir, AUDIO_EXT):
        m = REC_RE.match(p.stem)
        if not m or not 1 <= int(m[1]) <= 114:
            problems.append(f"recitation file not recognised: {p.name}")
        elif int(m[1]) in rec:
            problems.append(f"several recitations for surah {int(m[1])}; using {rec[int(m[1])].name}")
        else:
            rec[int(m[1])] = p
    if ref_dir:
        for p in _files(ref_dir, REF_EXT):
            m = TAF_RE.match(p.stem)
            if m and 1 <= int(m[2]) <= 114:
                ref.setdefault((m[1].strip().lower(), int(m[2])), p)
            else:                                    # "100 سورة العاديات.docx": no language prefix, number only
                m = REC_RE.match(p.stem)
                if m and 1 <= int(m[1]) <= 114:
                    ref.setdefault((None, int(m[1])), p)
                else:
                    problems.append(f"reference text not recognised: {p.name}")
    jobs = []
    for (lang, num), (p, name) in sorted(taf.items(), key=lambda kv: (kv[0][1], kv[0][0])):
        if num not in rec:
            problems.append(f"no recitation found for surah {num} ({lang})")
            continue
        r = ref.get((lang.lower(), num)) or ref.get((None, num))
        if ref_dir and not r:
            problems.append(f"no reference text for surah {num} ({lang}) - Gemini will work alone")
        jobs.append(Job(num, lang, name or f"surah {num}", p, rec[num], r))
    return jobs, problems


# ------------------------------------------------------------------------- text handling ------- #
_GLYPHS = re.compile(r"[\uFD40-\uFDCF\uFE70-\uFEFF\u06DE\u06E9]")   # verse-number / ornament glyphs of the docx font
_TASHKEEL = re.compile(r"[\u0610-\u061A\u064B-\u065F\u0670\u06D6-\u06ED\u0640]")


def norm(text: str) -> list[str]:
    t = _GLYPHS.sub(" ", text or "")
    t = _TASHKEEL.sub("", unicodedata.normalize("NFKC", t))
    t = re.sub(r"[0-9\u0660-\u0669\u06F0-\u06F9]", " ", t)
    t = re.sub("[أإآٱ]", "ا", t).replace("ى", "ي").replace("ة", "ه").replace("ؤ", "و").replace("ئ", "ي")
    return re.sub(r"[^\w\s]", " ", t.lower()).replace("_", " ").split()


def skel(tok: str) -> str:
    """Spelling-proof key: Whisper writes العاديات, the Quran text العٰديٰت - drop the weak letters after the first."""
    return tok[:1] + re.sub("[اويء]", "", tok[1:])


def toks(text: str) -> list[str]:
    return [skel(t) for t in norm(text)]


def short(text: str, n=12) -> str:
    w = text.split()
    return text if len(w) <= 2 * n else " ".join(w[:n]) + " … " + " ".join(w[-n:])


def _is_red(run) -> bool:
    try:
        c = run.font.color.rgb
    except Exception:
        return False
    return c is not None and c[0] >= 170 and c[1] < 110 and c[2] < 110


def read_reference(path) -> list[list[str]]:
    """Reference file -> [[verse_text, interpretation_text], ...]. In the docx the verse is the red text."""
    path = Path(path)
    if path.suffix.lower() == ".txt":
        paras = [p.strip() for p in re.split(r"\n\s*\n", path.read_text("utf-8", "ignore")) if p.strip()]
        return [["", p] for p in paras]
    import docx
    document = docx.Document(str(path))
    blocks, cur, mode = [], None, None
    for para in document.paragraphs:
        for run in para.runs:
            text = run.text
            if _is_red(run):
                if mode != "a":
                    cur = ["", ""]
                    blocks.append(cur)
                    mode = "a"
                cur[0] += text
            elif text.strip() and cur is not None:
                mode = "i"
                cur[1] += text
        if mode == "i":
            cur[1] += " "
    if not blocks:                                   # no red text at all: one block per paragraph
        return [["", p.text.strip()] for p in document.paragraphs if p.text.strip()]
    for b in blocks:
        b[0], rng = _clean_verse(b[0])
        b[1] = _LEADIN.sub("", b[1]).strip()
        b.append(rng)                                 # (first verse, last verse) or None
    return blocks


_LEADIN = re.compile(r"(قال\s+تعالى|وقال\s+تعالى|قال\s+الله\s+تعالى)\s*[:：]?\s*$")


def _digits(text: str) -> list[int]:
    return [int("".join(str(unicodedata.digit(c)) for c in run))
            for run in re.findall(r"[0-9\u0660-\u0669\u06F0-\u06F9]+", text)]


def _glyph_digits(text: str) -> list[int]:
    """The docx font draws the numbers of the 'الآية N' label with glyphs U+FD50..FD59 (= digits 0..9)."""
    return [int("".join(str(ord(c) - 0xFD50) for c in run)) for run in re.findall(r"[\uFD50-\uFD59]+", text)]


def _clean_verse(red: str):
    """Red text is 'ﵥ<verse> ١ﵤ ﵝ<surah> الآية ﵑﵜ': -> (the verse itself, (first, last) verse numbers or None).
    A block that holds several verses ('من الآية ٨ الى الآية ٩') gives a range."""
    m = re.search(r"\uFD65(.*?)\uFD64", red, re.S) or re.search(r"[\uFD3E\uFD3F﴿](.*?)[\uFD3E\uFD3F﴾]", red, re.S)
    body = m.group(1) if m else re.split(r"الآية|الاية", red)[0]
    nums = _digits(body) or _glyph_digits(red)
    t = re.sub(r"[0-9\u0660-\u0669\u06F0-\u06F9]+", " ", _GLYPHS.sub(" ", body))
    return re.sub(r"\s+", " ", t).strip(), ((min(nums), max(nums)) if nums else None)


def plan_units(blocks, n):
    """-> (blocks, units). A unit is one verse, or several consecutive verses that the book explains together
    (e.g. 8-9). If the blocks' verse numbers tile 1..n exactly, they ARE the units; otherwise fall back to
    'one block per verse' (folding stray quotations)."""
    rng = [b[2] if len(b) > 2 else None for b in blocks]
    if rng and all(rng):
        pos = 1
        for a, b in rng:
            if a != pos or b < a:
                break
            pos = b + 1
        else:
            if pos == n + 1:
                return blocks, [tuple(r) for r in rng]
    fb = fit_blocks(blocks, n)
    return (fb, [(i, i) for i in range(1, n + 1)]) if fb else (None, None)


def unit_label(u) -> str:
    return str(u[0]) if u[0] == u[1] else f"{u[0]}-{u[1]}"


def fit_blocks(blocks, n):
    """Exactly n blocks: fold a stray red quotation (the smallest block) into the one before it."""
    blocks = [list(b) for b in blocks]
    while len(blocks) > n:
        k = min(range(1, len(blocks)), key=lambda i: len(norm(blocks[i][0])) + len(norm(blocks[i][1])))
        blocks[k - 1][1] += " " + blocks[k][0] + " " + blocks[k][1]
        del blocks[k]
    return blocks if len(blocks) == n else None


def align_blocks(texts, words):
    """First pass, no AI: where does each reference block start in the Whisper transcript?
    -> (word index per block, fraction of blocks that matched)."""
    hyp, hw = [], []
    for i, w in enumerate(words):
        for tok in toks(w.t):
            hyp.append(tok)
            hw.append(i)
    n, starts, ptr, hit = len(texts), [None] * len(texts), 0, 0
    for i, t in enumerate(texts):
        ref = toks(t)
        if not ref:
            continue
        win = hyp[ptr: ptr + 3 * len(ref) + 150]
        runs = [m for m in difflib.SequenceMatcher(None, ref, win, autojunk=False).get_matching_blocks()
                if m.size >= 2]
        head = [m for m in runs if m.a < max(8, 0.35 * len(ref))]
        if not head:
            ptr += len(ref)
            continue
        best = max(head, key=lambda m: (m.size, -m.b))
        pos = max(ptr, ptr + best.b - best.a)
        starts[i] = hw[min(pos, len(hw) - 1)]
        ptr = ptr + max(m.b + m.size for m in runs)
        hit += 1
    known = [(i, s) for i, s in enumerate(starts) if s is not None]
    if not known:
        return None, 0.0
    for i in range(n):                               # interpolate the blocks nothing matched
        if starts[i] is None:
            lo = max(((j, s) for j, s in known if j < i), default=(-1, 0))
            hi = min(((j, s) for j, s in known if j > i), default=(n, len(words) - 1))
            starts[i] = lo[1] + (hi[1] - lo[1]) * (i - lo[0]) // max(hi[0] - lo[0], 1)
    for i in range(1, n):                            # strictly increasing
        starts[i] = max(starts[i], starts[i - 1] + 1)
    return starts, hit / n


def _match(ref, hyp):
    if not ref or not hyp:
        return 0.0
    sm = difflib.SequenceMatcher(None, ref, hyp, autojunk=False)
    return sum(m.size for m in sm.get_matching_blocks()) / len(ref)


def score_span(text, words) -> float:
    """How well does what Whisper heard inside the span match the reference - at BOTH ends.
    A low head = starts in the wrong place; a low tail = words left out or pushed to the next verse."""
    ref, hyp = toks(text), [t for w in words for t in toks(w.t)]
    return round(min(_match(ref[:50], hyp[:100]), _match(ref[-50:], hyp[-100:])), 2)


# ---------------------------------------------------------------------------- audio ----------- #
def load_audio(path, rate=RATE):
    from pydub import AudioSegment
    seg = AudioSegment.from_file(str(path), parameters=["-ac", "1", "-ar", str(rate)])
    return seg.set_frame_rate(rate).set_channels(1).set_sample_width(2)      # force one format, whatever ffmpeg did


class Energy:
    """10 ms loudness curve, used to put every cut in a real pause instead of mid-word."""

    def __init__(self, seg):
        x = np.asarray(seg.get_array_of_samples(), dtype=np.float32)
        f = max(seg.frame_rate // 100, 1)
        n = len(x) // f
        rms = np.sqrt((x[: n * f].reshape(n, f) ** 2).mean(1)) + 1e-9
        self.db = 20 * np.log10(rms / np.percentile(rms, 99))
        self.total = len(seg)
        # "quiet" = a little above this recording's own noise floor (clean TTS and a reverberant mosque differ)
        self.thr = float(np.clip(np.percentile(self.db, 10) + 10, -48, -30))

    def quiet(self, a, b, thr=-34):
        seg = self.db[max(a, 0) // 10: max(b, 0) // 10 + 1]
        return len(seg) == 0 or seg.max() < thr

    def back(self, ms, lo):                       # walk earlier until the audio is quiet
        while ms > lo and not self.quiet(ms - 50, ms):
            ms -= 10
        return max(ms, lo)

    def forward(self, ms, hi):                    # walk later until the audio is quiet
        while ms < hi and not self.quiet(ms, ms + 50):
            ms += 10
        return min(ms, hi)


def find_gap(energy, lo, hi, near_lo, near_hi, min_ms=60):
    """Quiet stretches inside [lo, hi] ms -> the one that best matches the expected gap [near_lo, near_hi]."""
    a, b = max(lo, 0) // 10, min(hi, energy.total) // 10
    q = energy.db[a: b + 1] < energy.thr
    runs, i = [], 0
    while i < len(q):
        if q[i]:
            j = i
            while j + 1 < len(q) and q[j + 1]:
                j += 1
            if (j - i + 1) * 10 >= min_ms:
                runs.append(((a + i) * 10, (a + j + 1) * 10))
            i = j + 1
        else:
            i += 1
    if not runs:
        return None
    mid = (near_lo + near_hi) / 2

    def score(r):
        overlap = max(0, min(r[1], near_hi) - max(r[0], near_lo))
        return (r[1] - r[0]) + 2 * overlap - 0.5 * abs((r[0] + r[1]) / 2 - mid)
    return max(runs, key=score)


def boundary(energy, e_prev, s_next, snap):
    """Where one clip ends and the next begins, found in the AUDIO (not trusted from Whisper's word times).
    -> (end of the earlier clip, start of the later clip, warning or None)"""
    lo = max(min(e_prev, s_next) - snap, 0)
    hi = min(max(e_prev, s_next) + snap, energy.total)
    g = find_gap(energy, lo, hi, min(e_prev, s_next), max(e_prev, s_next))
    if g:
        w = g[1] - g[0]
        return int(g[0] + min(180, w // 2)), int(g[1] - min(120, w // 2)), None
    seg = energy.db[lo // 10: hi // 10 + 1]                    # no real silence: cut at the quietest point
    p = lo + int(np.argmin(np.convolve(seg, np.ones(5) / 5, "same"))) * 10
    bad = energy.db[min(p // 10, len(energy.db) - 1)] > energy.thr + 12
    return p, p, ("no pause found here - the cut may be inside speech" if bad else None)


def make_spans(words, starts, energy, total_ms, head_zero=False, snap=400):
    """Word index of each verse's first word -> ([start_ms, end_ms] per verse, end of the spoken lead-in, warnings).
    Every boundary is placed in a real silence between two verses, so one wrong Whisper timestamp can only
    move its own boundary - it can never shift the verses after it."""
    n = len(starts)
    S = [words[starts[k]].s for k in range(n)]
    E = [words[(starts[k + 1] - 1) if k + 1 < n else len(words) - 1].e for k in range(n)]
    begs, ends, warn = [0] * n, [0] * n, {}
    for k in range(n - 1):
        ends[k], begs[k + 1], w = boundary(energy, E[k], S[k + 1], snap)
        if w:
            warn[k] = warn[k + 1] = w
    lead_end = 0
    if starts[0] > 0:                                           # something is spoken before verse 1 (the intro)
        lead_end, begs[0], w = boundary(energy, words[starts[0] - 1].e, S[0], snap)
        if w:
            warn[0] = w
    else:
        begs[0] = max(energy.back(S[0] - 40, max(S[0] - snap, 0)) - 60, 0)
    g = find_gap(energy, E[-1] - 300, min(E[-1] + 1500, total_ms), E[-1], E[-1] + 300, min_ms=100)
    ends[-1] = int(g[0] + min(180, (g[1] - g[0]) // 2)) if g else min(total_ms, E[-1] + 700)
    out = []
    for k in range(n):
        st = 0 if (head_zero and k == 0) else begs[k]
        en = max(ends[k], st + 300)
        if k + 1 < n and begs[k + 1] > st + 300:
            en = min(en, begs[k + 1])
        out.append([int(st), int(min(en, total_ms))])
    return out, int(lead_end), warn


def peaks(path, work_dir) -> np.ndarray:
    """10 ms peak envelope for the review waveform (cached)."""
    cache = Path(work_dir) / "cache" / (hashlib.md5(f"{path}{Path(path).stat().st_mtime_ns}".encode()).hexdigest() + ".npy")
    if cache.exists():
        return np.load(cache)
    x = np.abs(np.asarray(load_audio(path, 8000).get_array_of_samples(), dtype=np.int32))
    x = x[: len(x) // 80 * 80].reshape(-1, 80).max(1).astype(np.float32)
    cache.parent.mkdir(parents=True, exist_ok=True)
    np.save(cache, x)
    return x


def waveform_b64(pk, t0, t1, others, cur, W=1000, H=110) -> str:
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (W, H), (22, 24, 28))
    d = ImageDraw.Draw(img)
    X = lambda t: int((t - t0) / (t1 - t0) * W)
    for s, e in others:
        d.rectangle((X(s), 0, X(e), H), fill=(46, 50, 58))
    d.rectangle((X(cur[0]), 0, X(cur[1]), H), fill=(24, 62, 70))
    seg = pk[int(t0 * 100): int(t1 * 100)]
    if len(seg):
        idx = np.linspace(0, len(seg), W + 1).astype(int)
        top = max(float(np.percentile(pk, 99.5)), 1.0)
        for x in range(W):
            amp = seg[idx[x]: max(idx[x + 1], idx[x] + 1)].max() / top
            h = min(1.0, amp) * (H / 2 - 4)
            d.line((x, H / 2 - h, x, H / 2 + h), fill=(150, 170, 190))
    t = int(t0) + 1
    while t < t1:
        d.line((X(t), H - 8, X(t), H), fill=(110, 110, 110))
        if t % 2 == 0:
            d.text((X(t) + 2, H - 20), f"{t}s", fill=(130, 130, 130))
        t += 1
    d.line((X(cur[0]), 0, X(cur[0]), H), fill=(60, 220, 120), width=3)
    d.line((X(cur[1]), 0, X(cur[1]), H), fill=(240, 90, 90), width=3)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return base64.b64encode(buf.getvalue()).decode()


# ------------------------------------------------------------------------------ Whisper ------- #
_WHISPER = {}


def transcribe(path, lang, cfg, log, stop=None, retranscribe=False) -> list[Word]:
    """Word-timed transcript from faster-whisper, cached next to the output."""
    path = Path(path)
    st = path.stat()
    key = hashlib.md5(f"{path.name}|{st.st_size}|{st.st_mtime_ns}|{cfg.whisper_model}|{lang}".encode()).hexdigest()[:14]
    cache = cfg.work_dir / "cache" / f"{key}.words.json"
    if cache.exists() and not retranscribe:
        return [Word(*w) for w in json.loads(cache.read_text("utf-8"))]
    from faster_whisper import WhisperModel
    if cfg.whisper_model not in _WHISPER:
        log(f"  loading Whisper '{cfg.whisper_model}' (first time downloads it)...")
        _WHISPER[cfg.whisper_model] = WhisperModel(cfg.whisper_model, device="auto", compute_type="auto")
    segs, info = _WHISPER[cfg.whisper_model].transcribe(
        str(path), language=lang, word_timestamps=True, beam_size=5,
        condition_on_previous_text=False, vad_filter=True)
    words, last = [], -1
    for seg in segs:
        if stop is not None and stop.is_set():
            raise RuntimeError("stopped")
        for w in seg.words or []:
            if w.word.strip():
                words.append(Word(int(w.start * 1000), int(w.end * 1000), w.word.strip()))
        pct = int(100 * seg.end / max(info.duration, 1))
        if pct // 20 > last:
            last = pct // 20
            log(f"  {path.name[:40]}: {pct}%")
    cfg.usage["audio_s"] += info.duration
    if len(words) > 10:                              # drop absurd words (a 10 s "word" is a glitch)
        med = float(np.median([w.e - w.s for w in words]))
        words = [w for w in words if w.e - w.s <= max(4000, 8 * med)]
    cache.parent.mkdir(parents=True, exist_ok=True)
    cache.write_text(json.dumps([[w.s, w.e, w.t] for w in words], ensure_ascii=False), "utf-8")
    return words


# ----------------------------------------------------------------------------------- Gemini ---- #
SYSTEM = ("You are a careful Arabic/Quran text-alignment assistant. You receive a word-timed speech-to-text "
          "transcript (it contains recognition errors) and must say at which word each verse begins. "
          "Answer with JSON only.")

RULES = {
    "rec": ("The recording is a sheikh reciting the Quran. Before the first verse he may say the isti'adha "
            "(أعوذ بالله من الشيطان الرجيم) and/or the basmala (بسم الله الرحمن الرحيم). These are NOT verses, "
            "so the first verse starts after them - except in Al-Fatiha, where the basmala IS verse 1 (only "
            "the isti'adha before it is not a verse). The same phrase may recur in different verses (e.g. "
            "refrains): use the verse count and order to decide, one start per verse."),
    "taf": ("The recording explains the Quran verse by verse. It opens with the speaker saying the surah number "
            "and name - that is NOT part of any verse. Each verse's explanation starts at the first word of the "
            "text that belongs to that verse (it often begins by quoting the verse). Every word must belong to "
            "exactly one verse: none left out, none pushed into the next verse's explanation."),
}


def ask_json(cfg, prompt, log=print):
    from google import genai
    from google.genai import types
    client = genai.Client(api_key=cfg.api_key, http_options=types.HttpOptions(timeout=cfg.timeout_s * 1000))
    last = None
    for i in range(3):
        t0 = time.time()
        try:
            log(f"  ... asking Gemini ({cfg.model}), attempt {i + 1}/3 (waits up to {cfg.timeout_s}s)")
            r = client.models.generate_content(model=cfg.model, contents=prompt, config={
                "response_mime_type": "application/json", "temperature": 0, "system_instruction": SYSTEM})
            u = getattr(r, "usage_metadata", None)
            cfg.usage["calls"] += 1
            cfg.usage["in"] += getattr(u, "prompt_token_count", 0) or 0
            cfg.usage["out"] += getattr(u, "candidates_token_count", 0) or 0
            log(f"  ... Gemini answered in {time.time() - t0:.0f}s")
            return json.loads(r.text)
        except Exception as e:
            last = e
            log(f"  ! Gemini attempt {i + 1} failed after {time.time() - t0:.0f}s: {str(e)[:300]}")
            time.sleep(2 * (i + 1))
    raise RuntimeError(f"Gemini request failed: {last}")


def _clock(ms):
    return f"{ms // 60000:02d}:{ms // 1000 % 60:02d}"


def render(words, lo, hi):
    lines = []
    for i in range(0, hi - lo, 12):
        chunk = words[lo + i: min(lo + i + 12, hi)]
        lines.append(f"@{_clock(chunk[0].s)} " + " ".join(f"[{i + j}]{w.t}" for j, w in enumerate(chunk)))
    return "\n".join(lines)


def ai_refine(cfg, kind, words, cand, texts, n, job, log):
    """Gemini checks (or, with no candidate, creates) the verse starts, a window of verses at a time.
    -> (starts, flagged verse numbers)"""
    if not cand and n > 80:
        raise RuntimeError("this surah is long: it needs the reference text to work in windows")
    wsize = cfg.window if cand else n
    out, flagged = [], set()
    for a in range(0, n, wsize):
        b = min(n, a + wsize)
        lo = out[a - 1] if a > 0 else 0
        hi = min(len(words), cand[b] + 40) if (cand and b < n) else len(words)
        hint = [c - lo for c in cand[a:b]] if cand else None
        units = job.units or [(i, i) for i in range(1, n + 1)]
        ref = "\n".join(f"Item {k + 1} (verse {unit_label(units[k])}): {short(texts[k], 10)}"
                        for k in range(a, b)) if texts else "(none given)"
        merged = [unit_label(u) for u in units if u[0] != u[1]]
        prompt = (f"Surah {job.number} ({job.name}), {n} items in total.\n{RULES[kind]}\n\n"
                  + (f"Some items hold several consecutive verses that must be treated as ONE item (verses "
                     f"{', '.join(merged)}): mark only where the item begins; the verses inside it follow one after another.\n\n"
                     if merged else "")
                  + f"Mark items {a + 1} to {b} ({b - a} items). "
                  + (f"The transcript slice starts at the first word of item {a}; item {a + 1} starts later.\n"
                     if a > 0 else "The slice starts at the beginning of the recording.\n")
                  + f"\nReference text of each item (first and last words), from the book:\n{ref}\n\n"
                  + (f"A rough machine guess of the start indices (may be wrong): {hint}\n\n" if hint else "")
                  + f"Transcript, every word tagged [index]:\n{render(words, lo, hi)}\n\n"
                  f'Reply {{"verse_starts": [{b - a} strictly increasing integers, the [index] of the first '
                  f'word of each item {a + 1}..{b}], "notes": "<one short sentence>"}}')
        got, complaint = None, ""
        for _ in range(3):
            vs = ask_json(cfg, prompt + complaint, log).get("verse_starts")
            if not isinstance(vs, list) or len(vs) != b - a:
                complaint = f"\n\nRejected: I need exactly {b - a} numbers."
            elif not all(isinstance(x, int) and 0 <= x < hi - lo for x in vs) or \
                    any(y <= x for x, y in zip(vs, vs[1:])) or (a > 0 and vs[0] < 1):
                complaint = "\n\nRejected: indices must be integers inside the transcript and strictly increasing."
            else:
                got = vs
                break
        if got is None:
            if not hint:
                raise RuntimeError("Gemini could not mark the verses")
            log(f"  ! verses {a + 1}-{b}: Gemini's answer was unusable, keeping the text-match guess")
            flagged |= set(range(a + 1, b + 1))
            got = hint
        elif hint:
            flagged |= {a + 1 + i for i, (g, h) in enumerate(zip(got, hint)) if abs(g - h) > cfg.disagree_words}
        out += [lo + g for g in got]
    return out, flagged


def _mark(kind, words, texts, n, job, cfg, log):
    """-> (starts, flagged set, ai_checked)"""
    cand, frac = align_blocks(texts, words) if texts else (None, 0.0)
    if texts and frac < 0.5:
        log(f"  ! only {frac:.0%} of the reference sections were found in the {kind} transcript")
    if cfg.api_key:
        try:
            s, f = ai_refine(cfg, kind, words, cand if frac >= 0.3 else None, texts, n, job, log)
            return s, f, True
        except Exception as e:
            if not cand:
                raise
            log(f"  ! Gemini check failed ({e}); using the text-match result")
            return cand, set(), False
    if not cand:
        raise RuntimeError("needs a reference text or a Gemini API key")
    return cand, set(), False


def intro_floor(words, job) -> int:
    """Index of the first word AFTER the spoken intro (surah number + name). 0 = intro not found."""
    name = [t for t in toks(job.name) if t not in toks("سورة")]
    head = [toks(w.t) for w in words[:30]]
    for i in range(len(head)):
        if name and any(t == name[-1] for t in head[i]):
            return i + 1
    return 0


def mark(kind, words, texts, n, job, cfg, log):
    """-> (starts, flagged set, ai_checked). For the interpretation, verse 1 can never begin before the
    spoken intro (surah number + name) has finished."""
    starts, flagged, ok = _mark(kind, words, texts, n, job, cfg, log)
    if kind == "taf":
        floor = intro_floor(words, job)
        if floor == 0:
            log("  ! could not find the surah name at the start of the interpretation - check verse 1 in Review")
            flagged = set(flagged) | {1}
        elif starts[0] < floor:
            log(f"  ! verse 1 was placed inside the intro; moved to word {floor} (after the surah name)")
            starts = list(starts)
            starts[0] = floor
            for i in range(1, len(starts)):
                starts[i] = max(starts[i], starts[i - 1] + 1)
    return starts, flagged, ok


# ------------------------------------------------------------------------------- a job -------- #
def safe(s):
    return re.sub(r'[\\/:*?"<>|]', " ", s).strip()


def out_path(cfg, job):
    return Path(cfg.final_dir) / f"{job.lang}_{job.number:03d} {safe(job.name)}_Combined.mp3"


def json_path(cfg, job):
    return cfg.work_dir / f"{job.tag}.verses.json"


def _entries(words, starts, spans, texts, flagged, cfg, units, warn=None):
    out, n = [], len(starts)
    for k in range(n):
        ws = words[starts[k]: starts[k + 1] if k + 1 < n else len(words)]
        sc = score_span(texts[k], ws) if texts else None
        why = []
        if (k + 1) in flagged:
            why.append("Gemini moved this verse")
        if warn and k in warn:
            why.append(warn[k])
        if sc is not None and sc < cfg.min_score:
            why.append(f"text match only {sc:.0%}")
        out.append({"verse": k + 1, "label": unit_label(units[k]), "start": round(spans[k][0] / 1000, 3), "end": round(spans[k][1] / 1000, 3),
                    "score": sc, "flag": "; ".join(why), "heard": short(" ".join(w.t for w in ws), 14),
                    "expected": short(texts[k], 14) if texts else ""})
    return out


def process_job(job, cfg, log, stop=None):
    total = VERSE_COUNTS[job.number - 1]
    blocks, units = None, [(i, i) for i in range(1, total + 1)]
    if job.ref:
        try:
            raw = read_reference(job.ref)
            blocks, u = plan_units(raw, total)
            if blocks is None:
                log(f"  ! reference has {len(raw)} sections but the surah has {total} verses and their numbers do not "
                    f"line up - ignoring it for the text match")
            else:
                units = u
                merged = [unit_label(x) for x in u if x[0] != x[1]]
                if merged:
                    log(f"  reference explains these verses together, treated as one: {', '.join(merged)}")
        except Exception as e:
            log(f"  ! could not read the reference ({e})")
    n = len(units)
    job.units = units
    rec_texts = [b[0] for b in blocks] if blocks and all(b[0].strip() for b in blocks) else None
    taf_texts = [b[1] for b in blocks] if blocks and all(b[1].strip() for b in blocks) else None

    log("  transcribing the recitation...")
    rw = transcribe(job.rec, "ar", cfg, log, stop)
    log("  transcribing the interpretation...")
    tw = transcribe(job.taf, lang_code(job.lang), cfg, log, stop)
    if len(rw) < n or len(tw) < n:
        raise RuntimeError(f"Whisper heard too few words (recitation {len(rw)}, interpretation {len(tw)})")

    log("  marking the recitation verses...")
    rs, rf, rok = mark("rec", rw, rec_texts, n, job, cfg, log)
    log("  marking the interpretation verses...")
    ts, tf, tok = mark("taf", tw, taf_texts, n, job, cfg, log)

    rec_audio, taf_audio = load_audio(job.rec), load_audio(job.taf)
    rspans, _, rwarn = make_spans(rw, rs, Energy(rec_audio), len(rec_audio), head_zero=True, snap=cfg.snap_ms)   # verse 1 keeps isti'adha/basmala
    tspans, lead_end, twarn = make_spans(tw, ts, Energy(taf_audio), len(taf_audio), snap=cfg.snap_ms)
    data = {"surah": job.number, "name": job.name, "lang": job.lang, "verses": n, "units": [list(x) for x in units], "taf_intro_end": round(lead_end / 1000, 3),
            "rec_file": str(job.rec), "taf_file": str(job.taf),
            "rec_total": round(len(rec_audio) / 1000, 3), "taf_total": round(len(taf_audio) / 1000, 3),
            "gaps": [cfg.gap_after_rec_ms, cfg.gap_after_taf_ms], "bitrate": cfg.bitrate,
            "output": str(out_path(cfg, job)), "ai_checked": rok and tok,
            "rec": _entries(rw, rs, rspans, rec_texts, rf, cfg, units, rwarn),
            "taf": _entries(tw, ts, tspans, taf_texts, tf, cfg, units, twarn)}
    if data["taf"][0]["start"] < 0.3:
        data["taf"][0]["flag"] = (data["taf"][0]["flag"] + "; " if data["taf"][0]["flag"] else "") + \
            "no intro (surah number + name) before verse 1"
    bad = [e["verse"] for k in ("rec", "taf") for e in data[k] if e["flag"]]
    data["status"] = "review" if bad or (cfg.api_key and not data["ai_checked"]) else "ok"
    save_json(json_path(cfg, job), data)
    if data["status"] == "ok":
        assemble_json(data, log)
    else:
        log(f"  ✗ NEEDS REVIEW - {len(set(bad))} verse(s) flagged -> open the Review tab")
    return data


def save_json(path, data):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=1), "utf-8")


def assemble_json(data, log=print):
    """intro (surah number + name) -> for each verse: recitation, pause, interpretation, pause -> mp3."""
    from pydub import AudioSegment
    rec, taf = load_audio(data["rec_file"]), load_audio(data["taf_file"])
    ga, gb = data["gaps"]
    pieces = []

    def clip(seg, s, e):
        c = seg[int(s * 1000): int(e * 1000)]
        return c.fade_in(8).fade_out(8) if len(c) > 40 else c

    def silence(ms):
        return AudioSegment.silent(ms, OUT_RATE).set_channels(1).set_sample_width(2)

    def norm_fmt(c):                                  # every piece in the exact same format before joining
        return c.set_frame_rate(OUT_RATE).set_channels(1).set_sample_width(2)

    intro_end = data.get("taf_intro_end", data["taf"][0]["start"])
    if intro_end > 0.3:
        pieces += [norm_fmt(clip(taf, 0, intro_end)), silence(gb)]
    for r, t in zip(data["rec"], data["taf"]):
        pieces += [norm_fmt(clip(rec, r["start"], r["end"])), silence(ga),
                   norm_fmt(clip(taf, t["start"], t["end"])), silence(gb)]
    out = silence(0)._spawn(b"".join(p.raw_data for p in pieces))      # all pieces share one format now
    Path(data["output"]).parent.mkdir(parents=True, exist_ok=True)
    out.export(data["output"], format="mp3", bitrate=data.get("bitrate", "64k"))
    log(f"  ✓ saved {Path(data['output']).name}  ({len(out) / 60000:.1f} min)")
    return data["output"]


def run_batch(cfg, jobs, log, progress, stop):
    done = []
    for i, job in enumerate(jobs):
        if stop.is_set():
            log("⏹ stopped")
            break
        progress(i, len(jobs))
        log(f"[{i + 1}/{len(jobs)}] {job.lang} - surah {job.number} {job.name}")
        if out_path(cfg, job).exists() and not cfg.overwrite:
            log("  already built (tick 'Overwrite' to redo)")
            continue
        try:
            done.append(process_job(job, cfg, log, stop))
        except Exception as e:
            log(f"  ✗ ERROR: {e}")
    progress(len(jobs), max(len(jobs), 1))
    return done
